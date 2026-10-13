"""Exporta el borrador Markdown a un DOCX A4 basado en el esquema FISI local.

Uso: python docs/build_word.py. Conserva intacta la plantilla de referencia.
"""

from __future__ import annotations

import re
import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs/Estructura de TESIS (1).docx"
SOURCE = ROOT / "docs/tesis.md"
DEFAULT_OUTPUT = ROOT / "docs/tesis-borrador-2026-10-13.docx"
TOC_TITLES = [
    "Capítulo 1. Planteamiento del problema",
    "Capítulo 2. Marco teórico",
    "Capítulo 3. Estado del arte y método de investigación",
    "Capítulo 4. Aporte teórico",
    "Capítulo 5. Aporte práctico",
    "Capítulo 6. Análisis de resultados",
    "Conclusiones",
    "Recomendaciones",
    "Trabajos futuros",
    "Referencias bibliográficas",
    "Anexos",
]
DEFAULT_TOC_PAGES = ROOT / "docs/qa/toc-pages.json"


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        tag = "w:" + side
        item = borders.find(qn(tag))
        if item is None:
            item = OxmlElement(tag)
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")


def add_inline(paragraph, raw: str) -> None:
    raw = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", raw)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", raw)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part.strip("*`") if part[0] in "*`" else part)
        if part.startswith("**"):
            run.bold = True
        elif part.startswith("*"):
            run.italic = True
        elif part.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for child in (begin, instruction, separate, text, end):
        run._r.append(child)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Actualizar índice en Word (Ctrl+A, F9)."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for child in (begin, instruction, separate, text, end):
        run._r.append(child)


def make_table(doc, rows: list[list[str]]) -> None:
    width = len(rows[0])
    compact = len(rows) > 40
    table = doc.add_table(rows=1, cols=width)
    table.autofit = True
    for idx, text in enumerate(rows[0]):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        add_inline(cell.paragraphs[0], text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E6EBF0")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for values in rows[1:]:
        cells = table.add_row().cells
        for idx, text in enumerate(values[:width]):
            cells[idx].text = ""
            add_inline(cells[idx].paragraphs[0], text)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        for cell in row.cells:
            set_cell_border(cell)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0 if compact else 2)
                if compact:
                    para.paragraph_format.line_spacing = 1
                for run in para.runs:
                    run.font.size = Pt(7.5 if compact else 8)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--toc-pages", type=Path, default=DEFAULT_TOC_PAGES)
    args = parser.parse_args()
    doc = Document(TEMPLATE)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[style_name] if style_name in doc.styles else doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        if style_name.startswith("Heading "):
            p_pr = style.element.get_or_add_pPr()
            outline = p_pr.find(qn("w:outlineLvl"))
            if outline is None:
                outline = OxmlElement("w:outlineLvl")
                p_pr.append(outline)
            outline.set(qn("w:val"), str(int(style_name[-1]) - 1))

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    cover = lines[: lines.index("---")]
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(title, cover[0].lstrip("# "))
    for line in cover[1:]:
        if not line.strip() or line.startswith("**Estado del documento:**"):
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, line)
    doc.add_page_break()

    in_code = False
    skipping_manual_index = False
    i = lines.index("---") + 1
    while i < len(lines):
        line = lines[i].rstrip()
        if line == "## Índice":
            doc.add_heading("Índice general", level=1)
            pages = {}
            if args.toc_pages.exists():
                import json

                pages = json.loads(args.toc_pages.read_text(encoding="utf-8"))
            for title in TOC_TITLES:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.tab_stops.add_tab_stop(Cm(14), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                run = p.add_run(title + "\t" + str(pages.get(title, "?")))
                run.font.size = Pt(10)
            skipping_manual_index = True
            i += 1
            continue
        if skipping_manual_index:
            if line == "---":
                skipping_manual_index = False
            i += 1
            continue
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            i += 1
            continue
        if line.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                make_table(doc, rows)
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading = line[level:].strip()
            if level == 2 and heading in ("Abstract", "Índice de tablas"):
                doc.add_page_break()
            if level == 2 and (heading.startswith("Capítulo ") or heading in ("Conclusiones", "Recomendaciones", "Trabajos futuros", "Referencias bibliográficas", "Anexos")):
                doc.add_page_break()
            doc.add_heading(heading, level=min(3, max(1, level - 1)))
            i += 1
            continue
        if line == "---" or not line.strip() or line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("- ") or re.match(r"^\d+\. ", line):
            numbered = bool(re.match(r"^\d+\. ", line))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            if numbered:
                add_inline(p, line)
            else:
                add_inline(p, "• " + line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line.removeprefix("> "))
        i += 1

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
